"""Deterministic Excel / Markdown export from project session state.

Does not call Docling, Gemini, discovery, extraction, or validation.
Formats existing ProjectSpecSection / SectionSubmittalView / drafts only.
"""

from __future__ import annotations

import re
from collections import Counter
from enum import Enum
from io import BytesIO
from typing import Any, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from src.submittal.models import SubmittalType
from src.submittal.packaging import DraftSubmittalRow
from src.submittal.project import (
    ProcessingStatus,
    ProjectSpecSection,
    section_label,
)
from src.submittal.view_models import RequirementView, SectionSubmittalView

ROW_KIND_EXTRACTED = "Extracted Requirement"
ROW_KIND_CREATED = "Created Submittal"

PRODUCTS_STATUS_INCLUDED = "Included"
PRODUCTS_STATUS_AVAILABLE = "Available / Not Reviewed"

REGISTER_HEADERS = [
    "Source File",
    "Spec Section",
    "Spec Title",
    "Source Clause",
    "Submittal Type",
    "Submittal Title",
    "Products",
    "Product Count",
    "Products Status",
    "Row Kind",
]


class ExportRowKind(str, Enum):
    EXTRACTED = ROW_KIND_EXTRACTED
    CREATED = ROW_KIND_CREATED


class ExportSubmittalRow(BaseModel):
    """One exported register row (read/format model only)."""

    source_file: str
    project_section_key: str
    spec_section: str
    spec_title: str | None = None
    source_clause: str | None = None
    submittal_type: str
    submittal_title: str

    products: list[str] = Field(default_factory=list)
    product_count: int = 0
    products_status: str

    row_kind: str
    processing_status: str = ProcessingStatus.COMPLETED.value

    requirement_id: str | None = None
    draft_id: str | None = None


class ProjectExportSummary(BaseModel):
    """Project-level counts for export cover sheets / Markdown intro."""

    uploaded_file_count: int = 0
    detected_section_count: int = 0
    processed_section_count: int = 0
    completed_count: int = 0
    no_submittals_count: int = 0
    failed_count: int = 0
    not_processed_count: int = 0
    exported_register_row_count: int = 0
    status_lines: list[str] = Field(default_factory=list)


class ProjectExportBundle(BaseModel):
    """Complete export payload built from session/project state."""

    rows: list[ExportSubmittalRow] = Field(default_factory=list)
    summary: ProjectExportSummary = Field(default_factory=ProjectExportSummary)


def format_submittal_type(value: SubmittalType | str) -> str:
    raw = value.value if isinstance(value, SubmittalType) else str(value)
    return raw.replace("_", " ").title()


def products_excel_cell(products: Sequence[str]) -> str:
    return "\n".join(products)


def products_markdown_cell(products: Sequence[str]) -> str:
    return "; ".join(products)


def sanitize_export_basename(project_name: str | None = None) -> str:
    """Build a safe download basename without hard-coding a section number."""
    if not project_name or not project_name.strip():
        return "submittal_register"
    cleaned = re.sub(r"[^\w\-]+", "_", project_name.strip(), flags=re.UNICODE)
    cleaned = cleaned.strip("_")
    return f"{cleaned}_submittal_register" if cleaned else "submittal_register"


def can_export_project(
    *,
    statuses: Mapping[str, str],
) -> bool:
    """True when at least one section finished processing (completed / no_submittals)."""
    return any(
        s
        in (
            ProcessingStatus.COMPLETED.value,
            ProcessingStatus.NO_SUBMITTALS.value,
        )
        for s in statuses.values()
    )


def _product_name_lookup(req: RequirementView) -> dict[str, str]:
    names: dict[str, str] = {}
    for group in req.product_groups:
        for product in group.products:
            names[product.product_id] = product.name
    return names


def _confirmed_product_names(
    req: RequirementView, confirmed_ids: Sequence[str]
) -> list[str]:
    lookup = _product_name_lookup(req)
    names: list[str] = []
    for pid in confirmed_ids:
        name = lookup.get(pid)
        if name:
            names.append(name)
    return names


def _coerce_draft(item: Any) -> DraftSubmittalRow:
    if isinstance(item, DraftSubmittalRow):
        return item
    return DraftSubmittalRow.model_validate(item)


def _coerce_view(item: Any) -> SectionSubmittalView | None:
    if item is None:
        return None
    if isinstance(item, SectionSubmittalView):
        return item
    if isinstance(item, dict):
        return SectionSubmittalView.model_validate(item)
    return None


def build_project_export(
    *,
    project_sections: Mapping[str, ProjectSpecSection],
    views_by_key: Mapping[str, Any],
    statuses: Mapping[str, str],
    confirmed_by_key: Mapping[str, Mapping[str, Sequence[str]]],
    drafts_by_key: Mapping[str, Sequence[Any]],
    document_order: Sequence[str] | None = None,
    uploaded_file_count: int | None = None,
) -> ProjectExportBundle:
    """Build export rows + summary from current project state (no I/O, no Gemini).

    Register rows come only from ``completed`` sections (requirements + drafts).
    ``no_submittals`` / ``failed`` appear in the summary only.
    Display filter is ignored — export always covers all successfully processed
    sections.
    """
    # Preserve upload / discovery order.
    if document_order:
        doc_rank = {did: index for index, did in enumerate(document_order)}
    else:
        seen: list[str] = []
        for psec in project_sections.values():
            if psec.document_id not in seen:
                seen.append(psec.document_id)
        doc_rank = {did: index for index, did in enumerate(seen)}

    ordered_sections = sorted(
        project_sections.values(),
        key=lambda s: (
            doc_rank.get(s.document_id, 10_000),
            s.start_page,
            s.section_number,
            s.key,
        ),
    )

    status_counter: Counter[str] = Counter()
    status_lines: list[str] = []
    rows: list[ExportSubmittalRow] = []

    for psec in ordered_sections:
        status = statuses.get(psec.key, ProcessingStatus.NOT_PROCESSED.value)
        status_counter[status] += 1
        label = section_label(psec)
        if status == ProcessingStatus.NO_SUBMITTALS.value:
            status_lines.append(f"{label} — no submittals")
        elif status == ProcessingStatus.FAILED.value:
            status_lines.append(f"{label} — failed")
        elif status == ProcessingStatus.COMPLETED.value:
            status_lines.append(f"{label} — completed")
        elif status == ProcessingStatus.PROCESSING.value:
            status_lines.append(f"{label} — processing")
        else:
            status_lines.append(f"{label} — not processed")

        if status != ProcessingStatus.COMPLETED.value:
            continue

        view = _coerce_view(views_by_key.get(psec.key))
        confirmed_map = confirmed_by_key.get(psec.key) or {}

        if view is not None:
            for req in view.requirements:
                confirmed_ids = list(confirmed_map.get(req.requirement_id) or [])
                if confirmed_ids:
                    products = _confirmed_product_names(req, confirmed_ids)
                    products_status = PRODUCTS_STATUS_INCLUDED
                else:
                    # AI suggestions alone are never exported as included products.
                    products = []
                    products_status = PRODUCTS_STATUS_AVAILABLE
                rows.append(
                    ExportSubmittalRow(
                        source_file=psec.filename,
                        project_section_key=psec.key,
                        spec_section=psec.section_number,
                        spec_title=psec.section_title,
                        source_clause=req.source_clause,
                        submittal_type=format_submittal_type(req.submittal_type),
                        submittal_title=req.title,
                        products=products,
                        product_count=len(products),
                        products_status=products_status,
                        row_kind=ROW_KIND_EXTRACTED,
                        processing_status=status,
                        requirement_id=req.requirement_id,
                    )
                )

        for item in drafts_by_key.get(psec.key) or []:
            draft = _coerce_draft(item)
            products = list(draft.product_names)
            rows.append(
                ExportSubmittalRow(
                    source_file=psec.filename,
                    project_section_key=psec.key,
                    spec_section=psec.section_number,
                    spec_title=psec.section_title,
                    source_clause=draft.source_clause,
                    submittal_type=format_submittal_type(draft.submittal_type),
                    submittal_title=draft.title,
                    products=products,
                    product_count=len(products),
                    products_status=PRODUCTS_STATUS_INCLUDED,
                    row_kind=ROW_KIND_CREATED,
                    processing_status=status,
                    requirement_id=draft.requirement_id,
                    draft_id=draft.id,
                )
            )

    if uploaded_file_count is not None:
        file_count = uploaded_file_count
    elif document_order is not None:
        file_count = len(document_order)
    else:
        file_count = len({p.document_id for p in project_sections.values()})

    processed = (
        status_counter[ProcessingStatus.COMPLETED.value]
        + status_counter[ProcessingStatus.NO_SUBMITTALS.value]
        + status_counter[ProcessingStatus.FAILED.value]
    )

    summary = ProjectExportSummary(
        uploaded_file_count=file_count,
        detected_section_count=len(project_sections),
        processed_section_count=processed,
        completed_count=status_counter[ProcessingStatus.COMPLETED.value],
        no_submittals_count=status_counter[ProcessingStatus.NO_SUBMITTALS.value],
        failed_count=status_counter[ProcessingStatus.FAILED.value],
        not_processed_count=status_counter[ProcessingStatus.NOT_PROCESSED.value]
        + status_counter[ProcessingStatus.PROCESSING.value],
        exported_register_row_count=len(rows),
        status_lines=status_lines,
    )
    return ProjectExportBundle(rows=rows, summary=summary)


def export_excel_bytes(bundle: ProjectExportBundle) -> bytes:
    """Build an in-memory .xlsx workbook from an export bundle."""
    wb = Workbook()

    # --- Submittal Register ---
    ws = wb.active
    ws.title = "Submittal Register"
    ws.append(REGISTER_HEADERS)
    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for row in bundle.rows:
        ws.append(
            [
                row.source_file,
                row.spec_section,
                row.spec_title or "",
                row.source_clause or "",
                row.submittal_type,
                row.submittal_title,
                products_excel_cell(row.products),
                row.product_count,
                row.products_status,
                row.row_kind,
            ]
        )

    wrap_cols = {6, 7}  # Submittal Title, Products (1-indexed)
    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for idx, cell in enumerate(row_cells, start=1):
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=idx in wrap_cols,
            )

    widths = {
        "A": 36,
        "B": 14,
        "C": 28,
        "D": 14,
        "E": 20,
        "F": 42,
        "G": 36,
        "H": 12,
        "I": 22,
        "J": 20,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "A2"
    if ws.max_row >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(REGISTER_HEADERS))}{max(ws.max_row, 1)}"

    # --- Project Summary ---
    summary_ws = wb.create_sheet("Project Summary")
    summary_ws.append(["Metric", "Value"])
    for cell in summary_ws[1]:
        cell.font = header_font
    s = bundle.summary
    summary_rows = [
        ("Uploaded PDF files", s.uploaded_file_count),
        ("Detected sections", s.detected_section_count),
        ("Processed sections", s.processed_section_count),
        ("Completed", s.completed_count),
        ("No submittals", s.no_submittals_count),
        ("Failed", s.failed_count),
        ("Not processed", s.not_processed_count),
        ("Exported register rows", s.exported_register_row_count),
    ]
    for metric, value in summary_rows:
        summary_ws.append([metric, value])
    summary_ws.column_dimensions["A"].width = 28
    summary_ws.column_dimensions["B"].width = 14
    summary_ws.freeze_panes = "A2"

    # --- Processing Status ---
    status_ws = wb.create_sheet("Processing Status")
    status_ws.append(["Section", "Note"])
    for cell in status_ws[1]:
        cell.font = header_font
    for line in s.status_lines:
        if " — " in line:
            left, _, right = line.partition(" — ")
            status_ws.append([left, right])
        else:
            status_ws.append([line, ""])
    status_ws.column_dimensions["A"].width = 48
    status_ws.column_dimensions["B"].width = 18
    status_ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _md_escape_cell(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", " ")


def export_markdown(bundle: ProjectExportBundle) -> str:
    """Build a Markdown document grouped by specification section."""
    s = bundle.summary
    lines: list[str] = [
        "# Submittal Register",
        "",
        "## Project Summary",
        "",
        f"- Uploaded PDF files: **{s.uploaded_file_count}**",
        f"- Detected sections: **{s.detected_section_count}**",
        f"- Processed sections: **{s.processed_section_count}**",
        f"- Completed: **{s.completed_count}**",
        f"- No submittals: **{s.no_submittals_count}**",
        f"- Failed: **{s.failed_count}**",
        f"- Not processed: **{s.not_processed_count}**",
        f"- Exported register rows: **{s.exported_register_row_count}**",
        "",
    ]

    if s.status_lines:
        lines.append("### Processing Status")
        lines.append("")
        for line in s.status_lines:
            lines.append(f"- {line}")
        lines.append("")

    lines.append("## Register")
    lines.append("")

    if not bundle.rows:
        lines.append("_No register rows to export from completed sections._")
        lines.append("")
        return "\n".join(lines)

    # Group by project_section_key preserving row order.
    grouped: dict[str, list[ExportSubmittalRow]] = {}
    for row in bundle.rows:
        grouped.setdefault(row.project_section_key, []).append(row)

    for _key, group in grouped.items():
        first = group[0]
        title = first.spec_title or "Untitled"
        lines.append(f"## {first.spec_section} — {title}")
        lines.append("")
        lines.append(f"_Source file: `{first.source_file}`_")
        lines.append("")
        lines.append(
            "| Source File | Spec Section | Type | Submittal Title | Products | "
            "Source Clause | Status | Row Kind |"
        )
        lines.append("|---|---|---|---|---|---|---|---|")
        for row in group:
            lines.append(
                "| "
                + " | ".join(
                    [
                        _md_escape_cell(row.source_file),
                        _md_escape_cell(row.spec_section),
                        _md_escape_cell(row.submittal_type),
                        _md_escape_cell(row.submittal_title),
                        _md_escape_cell(products_markdown_cell(row.products) or "—"),
                        _md_escape_cell(row.source_clause or "—"),
                        _md_escape_cell(row.products_status),
                        _md_escape_cell(row.row_kind),
                    ]
                )
                + " |"
            )
        lines.append("")

    return "\n".join(lines)


def export_markdown_bytes(bundle: ProjectExportBundle) -> bytes:
    return export_markdown(bundle).encode("utf-8")
