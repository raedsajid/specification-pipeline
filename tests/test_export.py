"""Step 24: Excel / Markdown export from project session state (no Gemini)."""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from src.submittal.export import (
    PRODUCTS_STATUS_AVAILABLE,
    PRODUCTS_STATUS_INCLUDED,
    ROW_KIND_CREATED,
    ROW_KIND_EXTRACTED,
    build_project_export,
    can_export_project,
    export_excel_bytes,
    export_markdown,
    sanitize_export_basename,
)
from src.submittal.models import SubmittalType
from src.submittal.packaging import (
    DraftSubmittalRow,
    PackagingMode,
    build_draft_submittals,
)
from src.submittal.project import ProcessingStatus, ProjectSpecSection
from src.submittal.view_models import (
    ProductGroupView,
    ProductView,
    RequirementView,
    SectionSubmittalView,
)


def _section(
    *,
    key: str,
    document_id: str,
    filename: str,
    section_number: str,
    title: str,
    start_page: int = 1,
) -> ProjectSpecSection:
    return ProjectSpecSection(
        key=key,
        document_id=document_id,
        filename=filename,
        section_id=f"sec_{section_number}",
        section_number=section_number,
        section_title=title,
        start_page=start_page,
        end_page=start_page + 1,
    )


def _req(
    *,
    requirement_id: str,
    title: str,
    spec_section_id: str,
    products: list[tuple[str, str]] | None = None,
    suggested: list[str] | None = None,
    clause: str = "1.04.A",
) -> RequirementView:
    products = products or []
    group = ProductGroupView(
        group_id="grp_1",
        name="Products",
        products=[
            ProductView(product_id=pid, name=name, suggested_for_requirement=pid in (suggested or []))
            for pid, name in products
        ],
    )
    return RequirementView(
        requirement_id=requirement_id,
        spec_section_id=spec_section_id,
        title=title,
        requirement_text=title,
        submittal_type=SubmittalType.PRODUCT_DATA,
        source_clause=clause,
        suggested_product_ids=list(suggested or []),
        suggested_product_count=len(suggested or []),
        available_product_count=len(products),
        product_groups=[group] if products else [],
    )


def _view(
    *,
    document_id: str,
    section_number: str,
    title: str,
    requirements: list[RequirementView],
) -> SectionSubmittalView:
    return SectionSubmittalView(
        document_id=document_id,
        spec_section_id=section_number,
        section_number=section_number,
        section_title=title,
        requirements=requirements,
        catalog_product_count=sum(r.available_product_count for r in requirements),
        product_group_count=1,
    )


def _two_file_fixture():
    s15891 = _section(
        key="doc_a::sec_15891",
        document_id="doc_a",
        filename="D021779-15891 - METAL DUCTWORK.pdf",
        section_number="15891",
        title="METAL DUCTWORK",
        start_page=1,
    )
    s15570 = _section(
        key="doc_b::sec_15570",
        document_id="doc_b",
        filename="D021779-15570 - BOILER ACCESSORIES.pdf",
        section_number="15570",
        title="BOILER ACCESSORIES",
        start_page=1,
    )
    # Same req_001 / prod_001 in both sections — must not collide in export.
    products = [("prod_001", "Duct Sealant"), ("prod_002", "Duct Cement & Primer")]
    r15891 = _req(
        requirement_id="req_001",
        title="Manufacturer's Product Data",
        spec_section_id="15891",
        products=products,
        suggested=["prod_001"],
        clause="1.04.A",
    )
    r15570 = _req(
        requirement_id="req_001",
        title="Manufacturer's Product Data",
        spec_section_id="15570",
        products=[("prod_001", "Boiler Trim"), ("prod_002", "O'Malley Gasket")],
        suggested=["prod_001"],
        clause="1.04.B",
    )
    sections = {s15891.key: s15891, s15570.key: s15570}
    views = {
        s15891.key: _view(
            document_id="doc_a",
            section_number="15891",
            title="METAL DUCTWORK",
            requirements=[r15891],
        ),
        s15570.key: _view(
            document_id="doc_b",
            section_number="15570",
            title="BOILER ACCESSORIES",
            requirements=[r15570],
        ),
    }
    return sections, views, s15891, s15570, r15891, r15570


def test_empty_project_cannot_export():
    assert can_export_project(statuses={}) is False
    assert (
        can_export_project(
            statuses={"k": ProcessingStatus.NOT_PROCESSED.value}
        )
        is False
    )
    assert (
        can_export_project(statuses={"k": ProcessingStatus.FAILED.value}) is False
    )


def test_can_export_completed_or_no_submittals():
    assert can_export_project(
        statuses={"a": ProcessingStatus.COMPLETED.value}
    )
    assert can_export_project(
        statuses={"a": ProcessingStatus.NO_SUBMITTALS.value}
    )


def test_unprocessed_and_failed_excluded_from_register():
    sections, views, s15891, s15570, *_ = _two_file_fixture()
    # Only 15891 completed; 15570 not processed; add a failed third.
    failed = _section(
        key="doc_c::sec_x",
        document_id="doc_c",
        filename="other.pdf",
        section_number="15512",
        title="PIPING INSULATION",
    )
    sections[failed.key] = failed
    statuses = {
        s15891.key: ProcessingStatus.COMPLETED.value,
        s15570.key: ProcessingStatus.NOT_PROCESSED.value,
        failed.key: ProcessingStatus.FAILED.value,
    }
    bundle = build_project_export(
        project_sections=sections,
        views_by_key=views,
        statuses=statuses,
        confirmed_by_key={},
        drafts_by_key={},
        document_order=["doc_a", "doc_b", "doc_c"],
        uploaded_file_count=3,
    )
    assert all(r.spec_section == "15891" for r in bundle.rows)
    assert bundle.summary.uploaded_file_count == 3
    assert bundle.summary.detected_section_count == 3
    assert bundle.summary.completed_count == 1
    assert bundle.summary.failed_count == 1
    assert bundle.summary.not_processed_count == 1
    assert any(
        line.startswith("15512") and line.endswith("failed")
        for line in bundle.summary.status_lines
    )
    assert any("15570" in line and "not processed" in line for line in bundle.summary.status_lines)


def test_no_submittals_in_summary_not_register():
    s = _section(
        key="doc_a::sec_1",
        document_id="doc_a",
        filename="a.pdf",
        section_number="15570",
        title="BOILER ACCESSORIES",
    )
    bundle = build_project_export(
        project_sections={s.key: s},
        views_by_key={},
        statuses={s.key: ProcessingStatus.NO_SUBMITTALS.value},
        confirmed_by_key={},
        drafts_by_key={},
        uploaded_file_count=1,
    )
    assert bundle.rows == []
    assert bundle.summary.no_submittals_count == 1
    assert any("no submittals" in line for line in bundle.summary.status_lines)


def test_single_section_export():
    sections, views, s15891, s15570, *_ = _two_file_fixture()
    sections = {s15891.key: s15891}
    views = {s15891.key: views[s15891.key]}
    bundle = build_project_export(
        project_sections=sections,
        views_by_key=views,
        statuses={s15891.key: ProcessingStatus.COMPLETED.value},
        confirmed_by_key={},
        drafts_by_key={},
        uploaded_file_count=1,
    )
    assert len(bundle.rows) == 1
    assert bundle.rows[0].source_file.endswith("METAL DUCTWORK.pdf")
    assert bundle.rows[0].products_status == PRODUCTS_STATUS_AVAILABLE
    assert bundle.rows[0].products == []


def test_multi_file_export_distinct_identity():
    sections, views, s15891, s15570, *_ = _two_file_fixture()
    statuses = {
        s15891.key: ProcessingStatus.COMPLETED.value,
        s15570.key: ProcessingStatus.COMPLETED.value,
    }
    confirmed = {
        s15891.key: {"req_001": ["prod_001"]},
        s15570.key: {"req_001": ["prod_002"]},
    }
    bundle = build_project_export(
        project_sections=sections,
        views_by_key=views,
        statuses=statuses,
        confirmed_by_key=confirmed,
        drafts_by_key={},
        document_order=["doc_a", "doc_b"],
        uploaded_file_count=2,
    )
    assert len(bundle.rows) == 2
    files = {r.source_file for r in bundle.rows}
    assert len(files) == 2
    keys = {r.project_section_key for r in bundle.rows}
    assert keys == {s15891.key, s15570.key}
    # Same requirement_id, different products / sections
    by_sec = {r.spec_section: r for r in bundle.rows}
    assert by_sec["15891"].products == ["Duct Sealant"]
    assert by_sec["15570"].products == ["O'Malley Gasket"]
    assert by_sec["15891"].requirement_id == by_sec["15570"].requirement_id == "req_001"


def test_ai_suggestions_not_treated_as_confirmed():
    sections, views, s15891, *_ = _two_file_fixture()
    sections = {s15891.key: s15891}
    views = {s15891.key: views[s15891.key]}
    bundle = build_project_export(
        project_sections=sections,
        views_by_key=views,
        statuses={s15891.key: ProcessingStatus.COMPLETED.value},
        confirmed_by_key={},  # AI suggested prod_001 but user never confirmed
        drafts_by_key={},
    )
    row = bundle.rows[0]
    assert row.products_status == PRODUCTS_STATUS_AVAILABLE
    assert row.products == []


def test_confirmed_products_included():
    sections, views, s15891, *_ = _two_file_fixture()
    sections = {s15891.key: s15891}
    views = {s15891.key: views[s15891.key]}
    bundle = build_project_export(
        project_sections=sections,
        views_by_key=views,
        statuses={s15891.key: ProcessingStatus.COMPLETED.value},
        confirmed_by_key={s15891.key: {"req_001": ["prod_001", "prod_002"]}},
        drafts_by_key={},
    )
    row = bundle.rows[0]
    assert row.products_status == PRODUCTS_STATUS_INCLUDED
    assert row.products == ["Duct Sealant", "Duct Cement & Primer"]
    assert row.product_count == 2


def test_combined_and_one_per_product_drafts():
    sections, views, s15891, *_ = _two_file_fixture()
    sections = {s15891.key: s15891}
    view = views[s15891.key]
    req = view.requirements[0]
    combined = build_draft_submittals(
        req,
        ["prod_001", "prod_002"],
        PackagingMode.COMBINED,
        project_section_key=s15891.key,
    )
    one_each = build_draft_submittals(
        req,
        ["prod_001", "prod_002"],
        PackagingMode.ONE_PER_PRODUCT,
        project_section_key=s15891.key,
    )
    assert len(combined) == 1
    assert len(one_each) == 2

    bundle = build_project_export(
        project_sections=sections,
        views_by_key={s15891.key: view},
        statuses={s15891.key: ProcessingStatus.COMPLETED.value},
        confirmed_by_key={s15891.key: {"req_001": ["prod_001", "prod_002"]}},
        drafts_by_key={s15891.key: combined + one_each},
    )
    created = [r for r in bundle.rows if r.row_kind == ROW_KIND_CREATED]
    extracted = [r for r in bundle.rows if r.row_kind == ROW_KIND_EXTRACTED]
    assert len(extracted) == 1
    assert len(created) == 3
    combined_row = next(r for r in created if r.product_count == 2)
    assert combined_row.products == ["Duct Sealant", "Duct Cement & Primer"]
    assert combined_row.products_status == PRODUCTS_STATUS_INCLUDED
    singles = [r for r in created if r.product_count == 1]
    assert len(singles) == 2


def test_excel_workbook_structure_and_row_count():
    sections, views, s15891, s15570, *_ = _two_file_fixture()
    bundle = build_project_export(
        project_sections=sections,
        views_by_key=views,
        statuses={
            s15891.key: ProcessingStatus.COMPLETED.value,
            s15570.key: ProcessingStatus.COMPLETED.value,
        },
        confirmed_by_key={
            s15891.key: {"req_001": ["prod_001"]},
            s15570.key: {},
        },
        drafts_by_key={},
        uploaded_file_count=2,
    )
    data = export_excel_bytes(bundle)
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == [
        "Submittal Register",
        "Project Summary",
        "Processing Status",
    ]
    reg = wb["Submittal Register"]
    # header + 2 data rows
    assert reg.max_row == 3
    assert reg["A1"].value == "Source File"
    assert reg["G2"].value  # products cell present for confirmed row
    # wrap / freeze / autofilter
    assert reg.freeze_panes == "A2"
    assert reg.auto_filter.ref is not None
    summary = wb["Project Summary"]
    metrics = {summary.cell(r, 1).value: summary.cell(r, 2).value for r in range(2, 10)}
    assert metrics["Uploaded PDF files"] == 2
    assert metrics["Exported register rows"] == 2


def test_markdown_section_headings_and_special_chars():
    sections, views, s15891, s15570, *_ = _two_file_fixture()
    bundle = build_project_export(
        project_sections=sections,
        views_by_key=views,
        statuses={
            s15891.key: ProcessingStatus.COMPLETED.value,
            s15570.key: ProcessingStatus.COMPLETED.value,
        },
        confirmed_by_key={
            s15891.key: {"req_001": ["prod_002"]},
            s15570.key: {"req_001": ["prod_002"]},
        },
        drafts_by_key={},
    )
    md = export_markdown(bundle)
    assert md.startswith("# Submittal Register")
    assert "## Project Summary" in md
    assert "## 15891 — METAL DUCTWORK" in md
    assert "## 15570 — BOILER ACCESSORIES" in md
    assert "Duct Cement & Primer" in md
    assert "O'Malley Gasket" in md
    assert "D021779-15891 - METAL DUCTWORK.pdf" in md


def test_sanitize_basename_no_hardcoded_section():
    assert sanitize_export_basename(None) == "submittal_register"
    assert sanitize_export_basename("") == "submittal_register"
    assert "15891" not in sanitize_export_basename("My Project!")
    assert sanitize_export_basename("Acme HVAC").endswith("submittal_register")


def test_export_causes_no_extraction_or_gemini(monkeypatch=None):
    """Export must not import/call extractor or pipeline run paths."""
    import src.submittal.export as export_mod

    source = open(export_mod.__file__, encoding="utf-8").read()
    assert "extract_submittals" not in source
    assert "run_submittal_extraction" not in source
    assert "from src.submittal.extractor" not in source
    assert "from src.submittal.pipeline" not in source
    assert "DoclingDocument" not in source

    sections, views, s15891, *_ = _two_file_fixture()
    bundle = build_project_export(
        project_sections={s15891.key: s15891},
        views_by_key={s15891.key: views[s15891.key]},
        statuses={s15891.key: ProcessingStatus.COMPLETED.value},
        confirmed_by_key={},
        drafts_by_key={},
    )
    # Pure formatting — repeated calls identical
    a = export_excel_bytes(bundle)
    b = export_excel_bytes(bundle)
    assert a == b
    assert export_markdown(bundle) == export_markdown(bundle)


def test_document_order_preserved():
    sections, views, s15891, s15570, *_ = _two_file_fixture()
    # Reverse document_order → 15570 first
    bundle = build_project_export(
        project_sections=sections,
        views_by_key=views,
        statuses={
            s15891.key: ProcessingStatus.COMPLETED.value,
            s15570.key: ProcessingStatus.COMPLETED.value,
        },
        confirmed_by_key={},
        drafts_by_key={},
        document_order=["doc_b", "doc_a"],
    )
    assert [r.spec_section for r in bundle.rows] == ["15570", "15891"]


def test_extracted_vs_created_row_kinds():
    sections, views, s15891, *_ = _two_file_fixture()
    draft = DraftSubmittalRow(
        id="draft_abc",
        requirement_id="req_001",
        spec_section_id="15891",
        project_section_key=s15891.key,
        submittal_type=SubmittalType.PRODUCT_DATA,
        title="Manufacturer's Product Data",
        product_ids=["prod_001"],
        product_names=["Duct Sealant"],
        packaging_mode=PackagingMode.COMBINED,
    )
    bundle = build_project_export(
        project_sections={s15891.key: s15891},
        views_by_key={s15891.key: views[s15891.key]},
        statuses={s15891.key: ProcessingStatus.COMPLETED.value},
        confirmed_by_key={},
        drafts_by_key={s15891.key: [draft]},
    )
    kinds = {r.row_kind for r in bundle.rows}
    assert kinds == {ROW_KIND_EXTRACTED, ROW_KIND_CREATED}
